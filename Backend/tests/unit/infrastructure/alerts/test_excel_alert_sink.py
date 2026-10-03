from __future__ import annotations

import asyncio
from datetime import UTC, datetime
from pathlib import Path

import openpyxl

from control_layer.domain.models.alert import Alert, AlertUserRef
from control_layer.infrastructure.alerts.excel_alert_sink import ExcelAlertSink

ALERT_FIELDS = (
    "id",
    "created_at",
    "call_id",
    "user",
    "status",
    "stage",
    "rule_id",
    "severity",
    "owasp",
    "reason",
    "evidence",
)


def _alert(**overrides) -> dict:
    base = {
        "id": "al_1",
        "created_at": "2026-10-03T10:41:40Z",
        "call_id": "c_000139",
        "user": {"sub": "anna.kowalska", "name": "Anna Kowalska", "role": "developer"},
        "status": "MASKED",
        "stage": "dlp",
        "rule_id": "pii_masking",
        "severity": "medium",
        "owasp": ["LLM02"],
        "reason": "3 email addresses masked",
        "evidence": ["t.lis@example.com"],
    }
    base.update(overrides)
    return base


async def test_creates_workbook_with_header_row(tmp_path: Path) -> None:
    path = tmp_path / "alerts.xlsx"
    sink = ExcelAlertSink(path)

    await sink.emit(_alert())

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == list(ALERT_FIELDS)


async def test_appends_rows_readable_back(tmp_path: Path) -> None:
    path = tmp_path / "alerts.xlsx"
    sink = ExcelAlertSink(path)

    await sink.emit(_alert(id="al_1"))
    await sink.emit(_alert(id="al_2", rule_id="secrets_detection"))

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert rows[0][0] == "al_1"
    assert rows[1][0] == "al_2"
    assert rows[1][6] == "secrets_detection"


async def test_reopens_existing_workbook_without_duplicating_header(tmp_path: Path) -> None:
    path = tmp_path / "alerts.xlsx"
    sink1 = ExcelAlertSink(path)
    await sink1.emit(_alert(id="al_1"))

    sink2 = ExcelAlertSink(path)
    await sink2.emit(_alert(id="al_2"))

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.max_row == 3
    header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    assert header == list(ALERT_FIELDS)


async def test_creates_parent_directory(tmp_path: Path) -> None:
    path = tmp_path / "nested" / "alerts.xlsx"
    sink = ExcelAlertSink(path)

    await sink.emit(_alert())

    assert path.exists()


async def test_concurrent_emits_are_serialized(tmp_path: Path) -> None:
    path = tmp_path / "alerts.xlsx"
    sink = ExcelAlertSink(path)

    await asyncio.gather(*(sink.emit(_alert(id=f"al_{i}")) for i in range(20)))

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.max_row == 21
    ids = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert sorted(ids) == sorted(f"al_{i}" for i in range(20))


async def test_accepts_real_domain_alert_model(tmp_path: Path) -> None:
    path = tmp_path / "alerts.xlsx"
    sink = ExcelAlertSink(path)

    alert = Alert(
        id="al_real",
        created_at=datetime(2026, 10, 3, 10, 41, 40, tzinfo=UTC),
        call_id="c_000139",
        user=AlertUserRef(sub="anna.kowalska", name="Anna Kowalska", role="developer"),
        status="MASKED",
        stage="dlp",
        rule_id="pii_masking",
        severity="medium",
        owasp=["LLM02"],
        reason="3 email addresses masked",
        evidence=["t.lis@example.com"],
    )

    await sink.emit(alert)

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    row = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert row[0] == "al_real"
    assert row[6] == "pii_masking"
    assert "anna.kowalska" in row[3]


async def test_clear_deletes_the_file_and_the_next_alert_recreates_the_header(
    tmp_path: Path,
) -> None:
    path = tmp_path / "alerts.xlsx"
    sink = ExcelAlertSink(path)
    await sink.emit(_alert())
    assert path.exists()

    await sink.clear()
    assert not path.exists()
    await sink.clear()

    await sink.emit(_alert(id="al_2"))
    rows = list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))
    assert rows[0] == ALERT_FIELDS
    assert len(rows) == 2
