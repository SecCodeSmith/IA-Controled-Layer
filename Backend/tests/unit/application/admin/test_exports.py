from __future__ import annotations

import csv
import io
import json

import openpyxl
import pytest

from control_layer.application.use_cases.admin.exports import (
    ExportAlertsUseCase,
    ExportAuditUseCase,
    ExportNotAvailableError,
)


class FakeAuditRepository:
    def __init__(self, rows: list[dict]) -> None:
        self._rows = rows

    async def export_rows(self) -> list[dict]:
        return self._rows


_ROW_A = {
    "call_id": "c_1",
    "timestamp": "2026-10-03T10:00:00Z",
    "kind": "chat",
    "target": "llm.complete",
    "decision": {"status": "ALLOWED", "stage": None, "rule_id": None},
    "tokens": {"prompt": 1, "completion": 1, "total": 2},
    "cost_usd": 0.0,
}
_ROW_B = {
    "call_id": "c_2",
    "timestamp": "2026-10-03T10:05:00Z",
    "kind": "tool_call",
    "target": "github.get_readme",
    "decision": {"status": "BLOCKED", "stage": "policy", "rule_id": "prompt_injection_signatures"},
    "tokens": {"prompt": 0, "completion": 0, "total": 0},
    "cost_usd": 0.0,
}


async def test_export_jsonl_produces_one_json_object_per_row() -> None:
    use_case = ExportAuditUseCase(FakeAuditRepository([_ROW_A, _ROW_B]))

    export_file = await use_case.execute("jsonl")

    lines = export_file.content.decode("utf-8").strip().splitlines()
    assert len(lines) == 2
    assert json.loads(lines[0])["call_id"] == "c_1"
    assert json.loads(lines[1])["call_id"] == "c_2"
    assert export_file.media_type == "application/jsonl"
    assert export_file.filename.endswith(".jsonl")


async def test_export_csv_has_a_stable_header_and_rows() -> None:
    use_case = ExportAuditUseCase(FakeAuditRepository([_ROW_A, _ROW_B]))

    export_file = await use_case.execute("csv")

    reader = csv.DictReader(io.StringIO(export_file.content.decode("utf-8")))
    rows = list(reader)
    assert reader.fieldnames is not None
    assert "call_id" in reader.fieldnames
    assert rows[0]["call_id"] == "c_1"
    assert rows[1]["call_id"] == "c_2"
    assert export_file.media_type == "text/csv"
    assert export_file.filename.endswith(".csv")


async def test_export_xlsx_is_a_parseable_workbook() -> None:
    use_case = ExportAuditUseCase(FakeAuditRepository([_ROW_A, _ROW_B]))

    export_file = await use_case.execute("xlsx")

    workbook = openpyxl.load_workbook(io.BytesIO(export_file.content))
    sheet = workbook.active
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "call_id" in header
    first_data_row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    assert first_data_row[header.index("call_id")] == "c_1"
    assert (
        export_file.media_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert export_file.filename.endswith(".xlsx")


async def test_export_empty_audit_still_produces_valid_files() -> None:
    use_case = ExportAuditUseCase(FakeAuditRepository([]))

    jsonl_file = await use_case.execute("jsonl")
    csv_file = await use_case.execute("csv")
    xlsx_file = await use_case.execute("xlsx")

    assert jsonl_file.content == b""
    assert csv_file.content == b""
    workbook = openpyxl.load_workbook(io.BytesIO(xlsx_file.content))
    assert workbook.active.max_row == 1


class FakeAlertsPath:
    def __init__(self, tmp_path, exists: bool) -> None:
        self.path = tmp_path / "alerts.xlsx"
        if exists:
            workbook = openpyxl.Workbook()
            workbook.active.append(["id", "reason"])
            workbook.active.append(["al_1", "masked"])
            workbook.save(self.path)


async def test_export_alerts_returns_the_excel_file_bytes(tmp_path) -> None:
    fixture = FakeAlertsPath(tmp_path, exists=True)
    use_case = ExportAlertsUseCase(fixture.path)

    export_file = await use_case.execute()

    workbook = openpyxl.load_workbook(io.BytesIO(export_file.content))
    assert workbook.active["A1"].value == "id"
    assert export_file.filename == "alerts.xlsx"
    assert (
        export_file.media_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


async def test_export_alerts_raises_when_file_is_missing(tmp_path) -> None:
    fixture = FakeAlertsPath(tmp_path, exists=False)
    use_case = ExportAlertsUseCase(fixture.path)

    with pytest.raises(ExportNotAvailableError):
        await use_case.execute()
