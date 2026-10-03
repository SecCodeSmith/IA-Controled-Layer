from __future__ import annotations

import io

import httpx
import openpyxl
import pytest
import yaml

from tests.conftest import ADMIN_HEADERS, call_tool, chat, get_token, read_sse

pytestmark = [pytest.mark.integration, pytest.mark.asyncio(loop_scope="session")]


async def _generate_traffic(api: httpx.AsyncClient) -> dict[str, str]:
    anna = await get_token(api, "anna.kowalska")
    ids: dict[str, str] = {}
    allowed = await call_tool(api, anna, "ci", "get_run", {"pipeline": "e2e-login", "date": "d"})
    ids["allowed"] = allowed.json()["call_id"]
    masked = await call_tool(api, anna, "logs-db", "query", {"service": "auth", "since": "24h"})
    ids["masked"] = masked.json()["call_id"]
    blocked = await call_tool(api, anna, "hr-db", "find_approver", {"request": "x"})
    ids["blocked"] = blocked.json()["error"]["call_id"]
    escalated = await call_tool(
        api, anna, "github", "delete_branch", {"repo": "web-app", "branch": "feature/old-login"}
    )
    ids["escalated"] = escalated.json()["call_id"]
    return ids


async def test_admin_routes_require_token(api: httpx.AsyncClient) -> None:
    for path in ("/api/feed", "/api/audit", "/api/alerts", "/api/stats", "/api/metrics",
                 "/api/policy", "/api/reports/security", "/api/attack-suite/scenarios", "/metrics"):
        response = await api.get(path)
        assert response.status_code == 401, path
        assert response.json()["error"]["code"] == "identity_rejected"
    assert (await api.post("/api/demo/reset")).status_code == 401
    assert (await api.get("/api/stats", headers={"X-Admin-Token": "wrong"})).status_code == 401


async def test_admin_token_query_parameter_and_header_precedence(api: httpx.AsyncClient) -> None:
    ok = await api.get("/api/stats", params={"admin_token": "admin-dev-token"})
    assert ok.status_code == 200
    bad_query = await api.get("/api/stats", params={"admin_token": "nope"})
    assert bad_query.status_code == 401
    header_wins = await api.get(
        "/api/stats",
        params={"admin_token": "admin-dev-token"},
        headers={"X-Admin-Token": "wrong"},
    )
    assert header_wins.status_code == 401
    header_ok = await api.get(
        "/api/stats", params={"admin_token": "nope"}, headers=ADMIN_HEADERS
    )
    assert header_ok.status_code == 200


async def test_feed_and_audit_lists(api: httpx.AsyncClient) -> None:
    ids = await _generate_traffic(api)
    feed = await api.get("/api/feed", headers=ADMIN_HEADERS)
    assert feed.status_code == 200
    items = feed.json()["items"]
    assert {i["call_id"] for i in items} >= set(ids.values())
    row = next(i for i in items if i["call_id"] == ids["masked"])
    assert row["status"] == "MASKED"
    assert row["stage"] == "dlp"
    assert row["rule_id"] == "pii_masking"
    assert row["kind"] == "tool_call"
    assert row["target"] == "logs-db.query"
    assert row["user"] == {"sub": "anna.kowalska", "name": "Anna Kowalska", "role": "developer"}

    blocked_only = await api.get("/api/feed", params={"status": "BLOCKED"}, headers=ADMIN_HEADERS)
    assert {i["status"] for i in blocked_only.json()["items"]} == {"BLOCKED"}

    audit = await api.get(
        "/api/audit", params={"user": "anna.kowalska", "kind": "tool_call"}, headers=ADMIN_HEADERS
    )
    assert audit.status_code == 200
    assert all("proxy_latency_ms" in i for i in audit.json()["items"])


async def test_call_detail_shows_raw_vs_delivered_and_stage_timings(
    api: httpx.AsyncClient,
) -> None:
    ids = await _generate_traffic(api)
    response = await api.get(f"/api/audit/{ids['masked']}", headers=ADMIN_HEADERS)
    assert response.status_code == 200
    detail = response.json()
    assert detail["decision"]["status"] == "MASKED"
    assert detail["decision"]["rule_id"] == "pii_masking"
    assert detail["mcp_server"] == "logs-db"
    assert detail["identity"]["agent_id"] == "agent-anna-dev-7f3a"
    assert detail["items_masked"] == 3
    assert "t.lis@example.com" in detail["response"]["raw"]
    assert "t.lis@example.com" not in detail["response"]["delivered"]
    assert "[EMAIL_1]" in detail["response"]["delivered"]
    assert detail["response"]["raw"] != detail["response"]["delivered"]
    assert "pii_masking" in (detail["matched_rule_yaml"] or "")
    assert set(detail["latency"]["stages"]) == {
        "identity", "authorization", "dlp", "policy", "behavior", "resource", "audit"
    }

    missing = await api.get("/api/audit/c_999999", headers=ADMIN_HEADERS)
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "not_found"


async def test_audit_exports(api: httpx.AsyncClient) -> None:
    await _generate_traffic(api)
    jsonl = await api.get("/api/audit/export", params={"format": "jsonl"}, headers=ADMIN_HEADERS)
    assert jsonl.status_code == 200
    assert "attachment" in jsonl.headers["content-disposition"]
    assert len([line for line in jsonl.text.splitlines() if line.strip()]) >= 4

    csv = await api.get("/api/audit/export", params={"format": "csv"}, headers=ADMIN_HEADERS)
    assert csv.status_code == 200
    assert csv.text.splitlines()[0].startswith("call_id")

    xlsx = await api.get("/api/audit/export", params={"format": "xlsx"}, headers=ADMIN_HEADERS)
    assert xlsx.status_code == 200
    sheet = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet.max_row >= 5

    bad = await api.get("/api/audit/export", params={"format": "pdf"}, headers=ADMIN_HEADERS)
    assert bad.status_code == 422


async def test_alerts_list_and_excel_export(api: httpx.AsyncClient) -> None:
    await _generate_traffic(api)
    response = await api.get("/api/alerts", headers=ADMIN_HEADERS)
    assert response.status_code == 200
    items = response.json()["items"]
    rules = {i["rule_id"] for i in items}
    assert {"pii_masking", "role_provisioning", "destructive_requires_approval"} <= rules
    sample = next(i for i in items if i["rule_id"] == "role_provisioning")
    assert sample["status"] == "BLOCKED"
    assert sample["stage"] == "authorization"
    assert sample["user"]["sub"] == "anna.kowalska"

    filtered = await api.get(
        "/api/alerts", params={"rule_id": "pii_masking"}, headers=ADMIN_HEADERS
    )
    assert {i["rule_id"] for i in filtered.json()["items"]} == {"pii_masking"}

    export = await api.get("/api/alerts/export", headers=ADMIN_HEADERS)
    assert export.status_code == 200
    assert "alerts.xlsx" in export.headers["content-disposition"]
    sheet = openpyxl.load_workbook(io.BytesIO(export.content)).active
    header = [cell.value for cell in sheet[1]]
    assert "rule_id" in header
    assert sheet.max_row >= 4


async def test_stats_counts(api: httpx.AsyncClient) -> None:
    await _generate_traffic(api)
    stats = (await api.get("/api/stats", headers=ADMIN_HEADERS)).json()
    assert stats["total_calls"] == 4
    assert stats["allowed"] == 1
    assert stats["masked"] == 1
    assert stats["blocked"] == 1
    assert stats["escalated"] == 1
    assert stats["by_rule"]["pii_masking"] == 1
    assert stats["by_stage"]["authorization"] >= 1
    assert stats["by_role"]["developer"] == 4
    assert stats["policy"] == {"version": 3, "status": "LOADED"}
    assert stats["provider"]["name"] == "mock"
    assert 0 <= stats["posture_score"] <= 100
    assert any(r["sub"] == "anna.kowalska" for r in stats["risk"])


async def test_metrics_json_and_prometheus(api: httpx.AsyncClient) -> None:
    await _generate_traffic(api)
    metrics = (await api.get("/api/metrics", headers=ADMIN_HEADERS)).json()
    assert metrics["proxy"]["p95_ms"] >= metrics["proxy"]["p50_ms"] >= 0
    assert "dlp" in metrics["stages"]
    assert metrics["stages"]["dlp"]["count"] >= 1
    assert set(metrics["cache"]) == {"hits", "misses", "hit_ratio"}
    prom = await api.get("/metrics", headers=ADMIN_HEADERS)
    assert prom.status_code == 200
    assert prom.headers["content-type"].startswith("text/plain")
    assert "control_layer_calls_total 4" in prom.text


async def test_policy_view(api: httpx.AsyncClient, policy_text: str) -> None:
    response = await api.get("/api/policy", headers=ADMIN_HEADERS)
    assert response.status_code == 200
    body = response.json()
    assert body["version"] == 3
    assert body["status"] == "LOADED"
    assert body["error"] is None
    assert "version: 3" in body["raw_yaml"]
    expected_rules = len(yaml.safe_load(policy_text)["rules"])
    grouped = body["rules_by_stage"]
    assert sum(len(v) for v in grouped.values()) == expected_rules
    assert {"authorization", "dlp", "policy", "behavior"} <= set(grouped)
    assert any(r["id"] == "pii_masking" for r in grouped["dlp"])
    assert any(r["id"] == "role_provisioning" for r in grouped["authorization"])

    reloaded = await api.post("/api/policy/reload", headers=ADMIN_HEADERS)
    assert reloaded.status_code == 200
    assert reloaded.json()["status"] == "LOADED"


async def test_security_report(api: httpx.AsyncClient) -> None:
    await _generate_traffic(api)
    response = await api.get(
        "/api/reports/security", params={"period": "24h"}, headers=ADMIN_HEADERS
    )
    assert response.status_code == 200
    body = response.json()
    assert body["period"] == "24h"
    assert body["markdown"].lstrip().startswith("#")
    assert body["top_rules"]
    assert body["owasp_coverage"]
    assert isinstance(body["recommendations"], list)


async def test_demo_reset_clears_audit_and_alerts(api: httpx.AsyncClient) -> None:
    await _generate_traffic(api)
    assert (await api.get("/api/feed", headers=ADMIN_HEADERS)).json()["items"]
    reset = await api.post("/api/demo/reset", headers=ADMIN_HEADERS)
    assert reset.status_code == 200
    assert reset.json() == {"ok": True}
    assert (await api.get("/api/feed", headers=ADMIN_HEADERS)).json()["items"] == []
    assert (await api.get("/api/alerts", headers=ADMIN_HEADERS)).json()["items"] == []
    stats = (await api.get("/api/stats", headers=ADMIN_HEADERS)).json()
    assert stats["total_calls"] == 0
    token = await get_token(api, "anna.kowalska")
    me = (await api.get("/v1/me", headers={"Authorization": f"Bearer {token}"})).json()
    assert me["budget"]["tokens_used"] == 0


async def test_chat_usage_counts_against_budget(api: httpx.AsyncClient) -> None:
    token = await get_token(api, "anna.kowalska")
    await chat(api, token, "Summarize our deployment checklist for tomorrow.")
    me = (await api.get("/v1/me", headers={"Authorization": f"Bearer {token}"})).json()
    assert me["budget"]["tokens_used"] > 0


async def test_feed_sse_delivers_new_call(shared_app, api: httpx.AsyncClient) -> None:
    token = await get_token(api, "anna.kowalska")

    async def trigger() -> None:
        await call_tool(api, token, "ci", "get_run", {"pipeline": "e2e-login", "date": "d"})

    events = await read_sse(
        shared_app.app,
        "/api/feed/stream?admin_token=admin-dev-token",
        until=lambda seen: any(name == "feed" for name, _ in seen),
        trigger=trigger,
    )
    feed_events = [data for name, data in events if name == "feed"]
    assert feed_events[0]["target"] == "ci.get_run"
    assert feed_events[0]["status"] == "ALLOWED"


async def test_sse_requires_admin_token(api: httpx.AsyncClient) -> None:
    response = await api.get("/api/feed/stream")
    assert response.status_code == 401
